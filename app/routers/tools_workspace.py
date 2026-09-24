"""Tools workspace API using dedicated tables in the existing MyOffice Supabase."""
from __future__ import annotations

import csv, hashlib, hmac, io, secrets, uuid, zipfile
from datetime import datetime, timedelta, timezone
from typing import Any, Literal, Optional
from xml.etree import ElementTree

from fastapi import APIRouter, Depends, File, Form, Header, HTTPException, UploadFile
from openpyxl import load_workbook
from pydantic import BaseModel, Field
from app.supabase_client import rows, supabase

router = APIRouter()
TABLE = {"accounts":"tools_workspace_accounts","sessions":"tools_workspace_sessions","employees":"tools_workspace_employees","tools":"tools_workspace_equipment","history":"tools_workspace_history","changes":"tools_workspace_changes","evidence":"tools_workspace_evidence","source_registers":"tools_workspace_source_registers","usage":"tools_workspace_usage","errors":"tools_workspace_errors","feedback":"tools_workspace_feedback","notification_reads":"tools_workspace_notification_reads"}
FEEDBACK_BUCKET = "tools-workspace-feedback"
EVIDENCE_BUCKET = "tools-workspace-evidence"
SOURCE_REGISTERS_BUCKET = "tools-workspace-source-registers"
_test_mode = False
_memory: dict[str,list[dict[str,Any]]] = {key:[] for key in TABLE}

def _now(): return datetime.now(timezone.utc).isoformat()
def _hash_password(password: str, salt: Optional[bytes]=None):
    salt=salt or secrets.token_bytes(16); digest=hashlib.scrypt(password.encode(),salt=salt,n=2**14,r=8,p=1)
    return salt.hex(),digest.hex()
def _token_hash(token: str): return hashlib.sha256(token.encode()).hexdigest()
def _all(kind: str, newest=False):
    if _test_mode:
        data=[dict(row) for row in _memory[kind]]; return list(reversed(data)) if newest else data
    query=supabase.table(TABLE[kind]).select("*")
    if newest: query=query.order({"history":"event_at","source_registers":"uploaded_at"}.get(kind,"created_at"),desc=True)
    return rows(query.execute())
def _find(kind: str, field: str, value: Any):
    if _test_mode: return next((dict(row) for row in _memory[kind] if row.get(field)==value),None)
    data=rows(supabase.table(TABLE[kind]).select("*").eq(field,value).limit(1).execute()); return data[0] if data else None
def _insert(kind: str, row: dict[str,Any]):
    if _test_mode: _memory[kind].append(dict(row)); return dict(row)
    data=rows(supabase.table(TABLE[kind]).insert(row).execute())
    if not data: raise HTTPException(500,"The record could not be saved.")
    return data[0]
def _update(kind: str, row_id: str, values: dict[str,Any]):
    if _test_mode:
        row=next((item for item in _memory[kind] if item["id"]==row_id),None)
        if not row: raise HTTPException(404,"Record not found.")
        row.update(values); return dict(row)
    data=rows(supabase.table(TABLE[kind]).update(values).eq("id",row_id).execute())
    if not data: raise HTTPException(404,"Record not found.")
    return data[0]
def _account_role(account:dict[str,Any]):
    return account.get("role") or ("admin" if account.get("can_issue") else "viewer")
def _public(account):
    role=_account_role(account)
    return {"id":account["id"],"name":account["name"],"username":account["username"],"role":role,"department":account.get("department"),"can_issue":role=="issuer","created_at":account["created_at"]}
def _session(authorization: Optional[str]=Header(default=None)):
    if not authorization or not authorization.startswith("Bearer "): raise HTTPException(401,"Sign in to continue.")
    session=_find("sessions","token_hash",_token_hash(authorization[7:].strip()))
    if not session or session.get("expires_at","")<=_now(): raise HTTPException(401,"This session is no longer valid.")
    account=_find("accounts","id",session["account_id"])
    if not account: raise HTTPException(401,"This session is no longer valid.")
    return account
def _issuer(account=Depends(_session)):
    if _account_role(account)!="issuer": raise HTTPException(403,"Only an Issuer account can issue, receive, transfer or extend equipment.")
    if not account.get("department"): raise HTTPException(403,"This Issuer account needs an assigned department.")
    return account
def _admin(account=Depends(_session)):
    if _account_role(account)!="admin": raise HTTPException(403,"An administrator account is required.")
    return account
def _operator(account=Depends(_session)):
    if _account_role(account) not in {"admin","issuer"}: raise HTTPException(403,"An administrator or Issuer account is required.")
    return account
def _ensure_managed_department(account:dict[str,Any],department:Optional[str]):
    if _account_role(account)=="issuer" and department!=account.get("department"):
        raise HTTPException(403,f"This Issuer may manage only {account.get('department')} records.")

def _tool_state(tool: dict[str,Any], actor: str) -> dict[str,Any]:
    state=dict(tool)
    state.setdefault("id",str(uuid.uuid4())); state.setdefault("status","available"); state.setdefault("condition","Good")
    state.setdefault("equipment_kind","other-equipment"); state.setdefault("department","Engineering"); state.setdefault("archived",False)
    state.setdefault("custody",None); state.setdefault("row_version",1); state.setdefault("created_by",actor); state.setdefault("created_at",_now()); state.setdefault("updated_at",_now())
    return state

def _apply_change(before: Optional[dict[str,Any]], after: dict[str,Any], action: str, actor: str, detail: str="", employee: Optional[dict[str,Any]]=None, idempotency_key: Optional[str]=None):
    """Apply one equipment mutation and its audit/change records atomically in Postgres."""
    after=_tool_state(after,actor); event={"id":str(uuid.uuid4()),"detail":detail,"event_at":_now(),"employee_id":employee.get("id") if employee else None,"employee_name":employee.get("name") if employee else None,"metadata":{}}
    if _test_mode:
        existing=next((row for row in _memory["tools"] if row["id"]==after["id"]),None)
        if existing: existing.clear(); existing.update(after); existing["row_version"]=int((before or {}).get("row_version",0))+1
        else: _memory["tools"].append(after)
        _memory["history"].append({"id":event["id"],"tool_id":after["id"],"tool_name":after["name"],"action":action,"detail":detail,"actor_name":actor,"employee_id":event["employee_id"],"employee_name":event["employee_name"],"event_at":event["event_at"],"metadata":{}})
        change={"id":str(uuid.uuid4()),"tool_id":after["id"],"action":action,"before_state":before or after,"after_state":dict(after),"actor_name":actor,"created_at":_now(),"undone_at":None,"undone_by":None}; _memory["changes"].append(change)
        return {"tool":dict(after),"change_id":change["id"]}
    payload={"idempotency_key":idempotency_key or str(uuid.uuid4()),"tool_id":before.get("id") if before else None,"expected_version":before.get("row_version") if before else None,"before_state":before or after,"after_state":after,"action":action,"actor_name":actor,"event":event}
    try:
        response=supabase.rpc("tools_workspace_apply_change",{"payload":payload}).execute(); data=getattr(response,"data",None) or {}
        return data.get("result",data)
    except Exception as exc:
        message=str(exc)
        if "STALE_VERSION" in message: raise HTTPException(409,"This record changed elsewhere. Refresh and try again.") from exc
        if "CONFLICT" in message or "duplicate" in message.lower(): raise HTTPException(409,"That register number already exists.") from exc
        raise

class RegisterAccount(BaseModel):
    name:str=Field(min_length=2,max_length=100); username:str=Field(min_length=3,max_length=254,pattern=r"^[^\s]+$"); password:str=Field(min_length=6,max_length=128); can_issue:bool=False
class Login(BaseModel): username:str; password:str
class AccountRoleUpdate(BaseModel):
    role:Literal["admin","issuer","viewer"]; department:Optional[str]=Field(default=None,max_length=100)
class EmployeeInput(BaseModel):
    employee_number:str=Field(min_length=1,max_length=60); name:str=Field(min_length=2,max_length=120); department:str=Field(min_length=1,max_length=100); job_title:Optional[str]=Field(default=None,max_length=100); supervisor_name:Optional[str]=Field(default=None,max_length=120)
class ToolInput(BaseModel):
    register_number:str=Field(min_length=1,max_length=80); name:str=Field(min_length=2,max_length=160); make_model:Optional[str]=Field(default=None,max_length=200); serial_number:Optional[str]=Field(default=None,max_length=120); category:Optional[str]=Field(default=None,max_length=100); equipment_kind:str=Field(default="other-equipment",max_length=80); storage_location:str=Field(min_length=1,max_length=240); department:str=Field(default="Engineering",max_length=100); section:Optional[str]=Field(default=None,max_length=100); condition:str=Field(default="Good",max_length=60); notes:Optional[str]=Field(default=None,max_length=2000); approval_ref:Optional[str]=Field(default=None,max_length=160); calibration:Optional[str]=Field(default=None,max_length=160); specifications:dict[str,str]=Field(default_factory=dict)
class ToolUpdate(BaseModel):
    register_number:Optional[str]=Field(default=None,min_length=1,max_length=80); name:Optional[str]=Field(default=None,min_length=2,max_length=160); make_model:Optional[str]=Field(default=None,max_length=200); serial_number:Optional[str]=Field(default=None,max_length=120); category:Optional[str]=Field(default=None,max_length=100); equipment_kind:Optional[str]=Field(default=None,max_length=80); storage_location:Optional[str]=Field(default=None,min_length=1,max_length=240); department:Optional[str]=Field(default=None,max_length=100); section:Optional[str]=Field(default=None,max_length=100); condition:Optional[str]=Field(default=None,max_length=60); notes:Optional[str]=Field(default=None,max_length=2000); approval_ref:Optional[str]=Field(default=None,max_length=160); calibration:Optional[str]=Field(default=None,max_length=160); specifications:Optional[dict[str,str]]=None
class IssueInput(BaseModel):
    employee_id:str; location:str=Field(min_length=1,max_length=240); expected_return_at:Optional[str]=None; job_reference:Optional[str]=Field(default=None,max_length=160); assigned_equipment:list[str]=Field(default_factory=list,max_length=30); notes:Optional[str]=Field(default=None,max_length=1000)
class ReturnInput(BaseModel):
    location:str=Field(min_length=1,max_length=240); condition:Literal["Good","Damaged","Missing parts"]="Good"; notes:Optional[str]=Field(default=None,max_length=1000)
class MovementInput(BaseModel):
    kind:Literal["issue","return","transfer","extend"]; employee_id:Optional[str]=None; employee_name:Optional[str]=None; department:Optional[str]=None; location:Optional[str]=Field(default=None,max_length=240); expected_return_at:Optional[str]=None; job_reference:Optional[str]=Field(default=None,max_length=160); assigned_equipment:list[str]=Field(default_factory=list,max_length=30); condition:Optional[str]=Field(default="Good",max_length=60); notes:Optional[str]=Field(default=None,max_length=2000); approval_ref:Optional[str]=Field(default=None,max_length=160); calibration:Optional[str]=Field(default=None,max_length=160); idempotency_key:Optional[str]=None
class ImportCommit(BaseModel):
    target:Literal["equipment","employees"]; rows:list[dict[str,Any]]=Field(max_length=500)
class UsageInput(BaseModel): event:str=Field(min_length=1,max_length=100); detail:Optional[str]=Field(default=None,max_length=500)
class ErrorInput(BaseModel): message:str=Field(min_length=1,max_length=2000); source:Optional[str]=Field(default=None,max_length=500); stack:Optional[str]=Field(default=None,max_length=8000)
class NotificationReadInput(BaseModel): keys:list[str]=Field(max_length=500)
class MarkReadyInput(BaseModel):
    resolution_note:str=Field(min_length=2,max_length=1000)

def _notification_key(tool:dict[str,Any],kind:str):
    marker=(tool.get("custody") or {}).get("expected_return_at") if kind=="overdue" else tool.get("row_version",1)
    return f"{kind}:{tool['id']}:{marker or 'current'}"

def _effective_status(tool:dict[str,Any], now:Optional[datetime]=None):
    """Derive overdue state from custody so alerts do not depend on a scheduler."""
    status=tool.get("status")
    if status!="issued" or tool.get("archived"): return status
    due=(tool.get("custody") or {}).get("expected_return_at")
    if not due: return status
    try:
        deadline=datetime.fromisoformat(str(due).replace("Z","+00:00"))
        if deadline.tzinfo is None: deadline=deadline.replace(tzinfo=timezone.utc)
        return "overdue" if deadline.astimezone(timezone.utc)<(now or datetime.now(timezone.utc)) else status
    except (TypeError,ValueError):
        return status

def _effective_tool(tool:dict[str,Any]):
    return {**tool,"status":_effective_status(tool)}

def _active_notifications():
    alerts=[]
    for stored_tool in _all("tools"):
        tool=_effective_tool(stored_tool); kind=tool.get("status")
        if tool.get("archived") or kind not in {"overdue","attention"}: continue
        alerts.append({"key":_notification_key(tool,kind),"kind":kind,"tool_id":tool["id"],"tool_name":tool.get("name","Equipment"),"department":tool.get("department","Engineering")})
    return alerts

def _new_session(account_id):
    token=secrets.token_urlsafe(32); _insert("sessions",{"id":str(uuid.uuid4()),"token_hash":_token_hash(token),"account_id":account_id,"created_at":_now(),"expires_at":(datetime.now(timezone.utc)+timedelta(days=7)).isoformat()}); return token
@router.post("/auth/register",status_code=201)
def register(body:RegisterAccount):
    username=body.username.strip().lower()
    if _find("accounts","username",username): raise HTTPException(409,"That username is already in use.")
    # The first account bootstraps administration. Every later self-registration is view-only;
    # an administrator explicitly assigns Issuer access and its department.
    role="admin" if not _all("accounts") else "viewer"
    salt,digest=_hash_password(body.password); account=_insert("accounts",{"id":str(uuid.uuid4()),"name":body.name.strip(),"username":username,"salt":salt,"password_hash":digest,"role":role,"department":None,"can_issue":False,"created_at":_now()})
    return {"token":_new_session(account["id"]),"account":_public(account)}
@router.post("/auth/login")
def login(body:Login):
    account=_find("accounts","username",body.username.strip().lower())
    if not account: raise HTTPException(401,"Username or password is incorrect.")
    _,candidate=_hash_password(body.password,bytes.fromhex(account["salt"]))
    if not hmac.compare_digest(candidate,account["password_hash"]): raise HTTPException(401,"Username or password is incorrect.")
    return {"token":_new_session(account["id"]),"account":_public(account)}
@router.get("/auth/me")
def me(account=Depends(_session)): return _public(account)
@router.get("/accounts")
def list_accounts(_=Depends(_admin)): return [_public(account) for account in _all("accounts")]
@router.patch("/accounts/{account_id}")
def update_account_role(account_id:str,body:AccountRoleUpdate,admin=Depends(_admin)):
    account=_find("accounts","id",account_id)
    if not account: raise HTTPException(404,"Account was not found.")
    department=(body.department or "").strip() or None
    if body.role=="issuer" and not department: raise HTTPException(422,"Choose the department this Issuer may manage.")
    if account_id==admin["id"] and body.role!="admin": raise HTTPException(409,"Assign another administrator before changing your own administrator role.")
    updated=_update("accounts",account_id,{"role":body.role,"department":department if body.role=="issuer" else None,"can_issue":body.role=="issuer"})
    return _public(updated)

@router.get("/notifications")
def notifications(account=Depends(_session)):
    read_rows=[row for row in _memory["notification_reads"] if row.get("account_id")==account["id"]] if _test_mode else rows(supabase.table(TABLE["notification_reads"]).select("alert_key").eq("account_id",account["id"]).execute())
    read={row["alert_key"] for row in read_rows}
    alerts=[{**alert,"read":alert["key"] in read} for alert in _active_notifications()]
    return {"alerts":alerts,"unread_count":sum(not alert["read"] for alert in alerts)}

@router.post("/notifications/read",status_code=202)
def read_notifications(body:NotificationReadInput,account=Depends(_session)):
    active={alert["key"] for alert in _active_notifications()}; accepted=list(dict.fromkeys(key for key in body.keys if key in active))
    for key in accepted:
        row={"account_id":account["id"],"alert_key":key,"read_at":_now()}
        if _test_mode:
            existing=next((item for item in _memory["notification_reads"] if item["account_id"]==account["id"] and item["alert_key"]==key),None)
            if existing: existing.update(row)
            else: _memory["notification_reads"].append(row)
        else: supabase.table(TABLE["notification_reads"]).upsert(row,on_conflict="account_id,alert_key").execute()
    return {"read":accepted}

@router.get("/employees")
def list_employees(_=Depends(_session)): return _all("employees")
@router.post("/employees",status_code=201)
def create_employee(body:EmployeeInput,account=Depends(_operator)):
    _ensure_managed_department(account,body.department)
    number=body.employee_number.strip()
    if any(row["employee_number"].lower()==number.lower() for row in _all("employees")): raise HTTPException(409,"That employee number already exists.")
    return _insert("employees",{"id":str(uuid.uuid4()),**body.model_dump(),"employee_number":number,"active":True,"created_at":_now(),"created_by":account["name"]})
@router.get("/tools")
def list_tools(_=Depends(_session)):
    tools=_all("tools")
    evidence=_all("evidence")
    by_tool:dict[str,list[dict[str,Any]]]={}
    for item in evidence:
        if item.get("removed_at"): continue
        row=dict(item)
        if not _test_mode:
            try:
                signed=supabase.storage.from_(EVIDENCE_BUCKET).create_signed_url(item["storage_path"],3600)
                row["url"]=signed.get("signedURL") or signed.get("signedUrl")
            except Exception: row["url"]=None
        by_tool.setdefault(item["tool_id"],[]).append(row)
    return [{**_effective_tool(tool),"evidence":by_tool.get(tool["id"],[])} for tool in tools]
@router.post("/tools",status_code=201)
def create_tool(body:ToolInput,account=Depends(_operator)):
    _ensure_managed_department(account,body.department)
    number=body.register_number.strip()
    if any(row["register_number"].lower()==number.lower() for row in _all("tools")): raise HTTPException(409,"That register number already exists.")
    result=_apply_change(None,{"id":str(uuid.uuid4()),**body.model_dump(),"register_number":number,"status":"available","custody":None},"created",account["name"],f"Added {number} to the register")
    return result["tool"]
@router.patch("/tools/{tool_id}")
def update_tool(tool_id:str,body:ToolUpdate,account=Depends(_operator)):
    tool=_find("tools","id",tool_id)
    if not tool: raise HTTPException(404,"Tool was not found.")
    _ensure_managed_department(account,tool.get("department"))
    values=body.model_dump(exclude_unset=True); after={**tool,**values}
    result=_apply_change(tool,after,"updated",account["name"],"Equipment details updated")
    # The additive specifications column post-dates the original atomic RPC. Keep the
    # privilege boundary in Python instead of broadening a SECURITY DEFINER function.
    if "specifications" in values:
        saved=_update("tools",tool_id,{"specifications":values["specifications"] or {}})
        _update("changes",result["change_id"],{"after_state":saved})
        return saved
    return result["tool"]
@router.post("/tools/{tool_id}/archive")
def archive_tool(tool_id:str,account=Depends(_operator)):
    tool=_find("tools","id",tool_id)
    if not tool: raise HTTPException(404,"Tool was not found.")
    _ensure_managed_department(account,tool.get("department"))
    if tool.get("custody"): raise HTTPException(409,"Receive this tool before archiving it.")
    after={**tool,"archived":not bool(tool.get("archived"))}
    action="restored" if after["archived"] is False else "archived"
    return _apply_change(tool,after,action,account["name"],f"Equipment {action}")["tool"]

@router.post("/tools/{tool_id}/mark-ready")
def mark_tool_ready(tool_id:str,body:MarkReadyInput,account=Depends(_operator)):
    tool=_find("tools","id",tool_id)
    if not tool: raise HTTPException(404,"Tool was not found.")
    _ensure_managed_department(account,tool.get("department"))
    if tool.get("archived"): raise HTTPException(409,"Restore this tool before marking it ready for use.")
    if tool.get("custody"): raise HTTPException(409,"Receive this tool before marking it ready for use.")
    if tool.get("status")!="attention": raise HTTPException(409,"Only equipment held for attention can be marked ready for use.")
    note=body.resolution_note.strip()
    after={**tool,"status":"available","condition":"Good","notes":note}
    return _apply_change(tool,after,"released",account["name"],f"Marked ready for use: {note}")["tool"]
@router.get("/history")
def history(_=Depends(_session)): return _all("history",True)

@router.post("/analytics/usage",status_code=202)
def capture_usage(body:UsageInput,account=Depends(_session)): return _insert("usage",{"id":str(uuid.uuid4()),"event":body.event,"detail":body.detail,"account_id":account["id"],"account_name":account["name"],"created_at":_now()})
@router.post("/analytics/errors",status_code=202)
def capture_error(body:ErrorInput,account=Depends(_session)): return _insert("errors",{"id":str(uuid.uuid4()),**body.model_dump(),"account_id":account["id"],"account_name":account["name"],"created_at":_now()})
@router.get("/analytics")
def analytics(_=Depends(_admin)):
    usage,errors,feedback=_all("usage",True),_all("errors",True),_all("feedback",True)
    if not _test_mode:
        for item in feedback:
            if item.get("audio_path"):
                try:
                    signed=supabase.storage.from_(FEEDBACK_BUCKET).create_signed_url(item["audio_path"],3600)
                    item["audio_url"]=signed.get("signedURL") or signed.get("signedUrl")
                except Exception:
                    item["audio_url"]=None
    return {"usage":usage,"errors":errors,"feedback":feedback,"totals":{"usage":len(usage),"errors":len(errors),"feedback":len(feedback)}}
@router.post("/feedback",status_code=201)
async def create_feedback(text:str=Form(default=""),audio:Optional[UploadFile]=File(default=None),account=Depends(_session)):
    if not text.strip() and audio is None: raise HTTPException(400,"Write a suggestion or attach an audio recording.")
    data=await audio.read() if audio else b""
    if len(data)>25*1024*1024: raise HTTPException(413,"Audio feedback must be smaller than 25 MB.")
    row_id,path=str(uuid.uuid4()),None
    content_type=(audio.content_type or "audio/webm").split(";",1)[0].lower() if audio else None
    if audio and content_type not in {"audio/webm","audio/ogg","audio/mp4","audio/mpeg","audio/wav","audio/x-wav"}: raise HTTPException(415,"Choose a WebM, OGG, MP4, MP3 or WAV audio recording.")
    if audio and not _test_mode:
        extensions={"audio/webm":"webm","audio/ogg":"ogg","audio/mp4":"m4a","audio/mpeg":"mp3","audio/wav":"wav","audio/x-wav":"wav"}; path=f"{account['id']}/{row_id}.{extensions[content_type]}"
        supabase.storage.from_(FEEDBACK_BUCKET).upload(path,data,{"content-type":content_type,"cache-control":"3600"})
    return _insert("feedback",{"id":row_id,"text":text.strip() or None,"audio_path":path,"audio_filename":audio.filename if audio else None,"audio_content_type":content_type,"audio_size":len(data),"account_id":account["id"],"account_name":account["name"],"created_at":_now()})

def _move(tool_id:str,body:MovementInput,account:dict[str,Any]):
    tool=_find("tools","id",tool_id)
    if not tool: raise HTTPException(404,"Tool was not found.")
    assigned_department=(account.get("department") or "").strip()
    if tool.get("department")!=assigned_department: raise HTTPException(403,f"This Issuer may record movements only for {assigned_department} equipment.")
    open_loan=tool.get("custody") and tool.get("status") in {"issued","overdue"}
    if body.kind=="issue" and (tool.get("archived") or tool.get("status")!="available"): raise HTTPException(409,"This tool is not available to issue.")
    if body.kind!="issue" and not open_loan: raise HTTPException(409,"This tool has no open loan.")
    employee=None
    if body.kind in {"issue","transfer"}:
        if body.employee_id: employee=_find("employees","id",body.employee_id) or _find("employees","employee_number",body.employee_id)
        if not employee and body.employee_name: employee=next((row for row in _all("employees") if row["name"].lower()==body.employee_name.lower()),None)
        if not employee: raise HTTPException(404,"Choose an employee from the Tools employee register.")
        if employee.get("department")!=assigned_department: raise HTTPException(403,f"Choose an employee from {assigned_department}.")
    now=_now(); after=dict(tool); location=(body.location or tool.get("storage_location") or "Unspecified").strip()
    if body.kind=="return":
        after.update(status="available" if body.condition=="Good" else "attention",storage_location=location,condition=body.condition or "Good",notes=body.notes,custody=None,calibration=body.calibration or tool.get("calibration"))
    elif body.kind=="extend":
        custody={**tool["custody"],"expected_return_at":body.expected_return_at,"notes":body.notes}; after.update(status="issued",custody=custody,notes=body.notes)
    else:
        old=tool.get("custody") or {}; custody={**old,"employee_id":employee["id"],"employee_number":employee["employee_number"],"employee_name":employee["name"],"department":employee["department"],"location":location,"job_reference":body.job_reference or old.get("job_reference"),"assigned_equipment":body.assigned_equipment or old.get("assigned_equipment",[]),"expected_return_at":body.expected_return_at or old.get("expected_return_at"),"issued_at":old.get("issued_at") or now,"original_due_at":old.get("original_due_at") or body.expected_return_at,"issued_by":account["name"],"notes":body.notes,"approval_ref":body.approval_ref}; after.update(status="issued",storage_location=location,custody=custody,notes=body.notes,approval_ref=body.approval_ref or tool.get("approval_ref"))
    targets=", ".join(body.assigned_equipment)
    detail={"issue":f"Issued to {employee['name']}"+(f" for {targets}" if targets else "") if employee else "Issued","return":f"Returned to {location}","transfer":f"Transferred to {employee['name']}" if employee else "Transferred","extend":f"Return date changed to {body.expected_return_at}"}[body.kind]
    result=_apply_change(tool,after,body.kind,account["name"],detail,employee,body.idempotency_key)
    return {"tool":result["tool"],"change_id":result["change_id"]}

@router.post("/tools/{tool_id}/commands")
def move_tool(tool_id:str,body:MovementInput,account=Depends(_issuer)): return _move(tool_id,body,account)
@router.post("/tools/{tool_id}/issue")
def issue_tool(tool_id:str,body:IssueInput,account=Depends(_issuer)):
    return _move(tool_id,MovementInput(kind="issue",employee_id=body.employee_id,location=body.location,expected_return_at=body.expected_return_at,job_reference=body.job_reference,assigned_equipment=body.assigned_equipment,notes=body.notes),account)
@router.post("/tools/{tool_id}/return")
def return_tool(tool_id:str,body:ReturnInput,account=Depends(_issuer)):
    return _move(tool_id,MovementInput(kind="return",location=body.location,condition=body.condition,notes=body.notes),account)

@router.post("/tools/{tool_id}/evidence",status_code=201)
async def upload_evidence(tool_id:str,files:list[UploadFile]=File(...),account=Depends(_operator)):
    tool=_find("tools","id",tool_id)
    if not tool: raise HTTPException(404,"Tool was not found.")
    _ensure_managed_department(account,tool.get("department"))
    allowed={"image/jpeg","image/png","image/webp","image/gif","image/avif","application/pdf"}; saved=[]
    for upload in files:
        content=await upload.read(); content_type=upload.content_type or "application/octet-stream"
        if content_type not in allowed: raise HTTPException(415,"Choose a JPG, PNG, WebP, GIF, AVIF or PDF file.")
        if not content: raise HTTPException(422,"An attachment was empty.")
        evidence_id=str(uuid.uuid4()); extension=(upload.filename or "file").rsplit(".",1)[-1]; path=f"{tool_id}/{evidence_id}.{extension}"
        if not _test_mode: supabase.storage.from_(EVIDENCE_BUCKET).upload(path,content,{"content-type":content_type})
        row={"id":evidence_id,"tool_id":tool_id,"storage_path":path,"original_name":upload.filename or "Attachment","content_type":content_type,"size_bytes":len(content),"uploaded_by":account["name"],"uploaded_at":_now(),"removed_at":None}
        try: saved.append(_insert("evidence",row))
        except Exception:
            if not _test_mode:
                try: supabase.storage.from_(EVIDENCE_BUCKET).remove([path])
                except Exception: pass
            raise
    _apply_change(tool,tool,"attachments",account["name"],f"Added {len(saved)} attachment(s)")
    return saved

@router.get("/source-registers")
def list_source_registers(_=Depends(_session)):
    saved=_all("source_registers",True)
    if not _test_mode:
        for item in saved:
            try:
                signed=supabase.storage.from_(SOURCE_REGISTERS_BUCKET).create_signed_url(item["storage_path"],3600)
                item["url"]=signed.get("signedURL") or signed.get("signedUrl")
            except Exception: item["url"]=None
    return saved

@router.post("/source-registers",status_code=201)
async def upload_source_register(department:str=Form(...),notes:str=Form(default=""),file:UploadFile=File(...),account=Depends(_session)):
    content=await file.read(); filename=file.filename or "Source register"; extension=filename.lower().rsplit(".",1)[-1] if "." in filename else ""
    allowed_extensions={"pdf","xlsx","xlsm","xls","csv","docx","doc","ods","odt","jpg","jpeg","png","webp"}
    if extension not in allowed_extensions: raise HTTPException(415,"Choose a PDF, Excel, CSV, Word, OpenDocument or image file.")
    if not content: raise HTTPException(422,"The source register was empty.")
    if len(content)>50*1024*1024: raise HTTPException(413,"Source registers must be smaller than 50 MB.")
    register_id=str(uuid.uuid4()); path=f"{datetime.now(timezone.utc):%Y/%m}/{register_id}.{extension}"; content_type=file.content_type or "application/octet-stream"
    if not _test_mode: supabase.storage.from_(SOURCE_REGISTERS_BUCKET).upload(path,content,{"content-type":content_type})
    row={"id":register_id,"department":department.strip() or "Unassigned","notes":notes.strip() or None,"storage_path":path,"original_name":filename,"content_type":content_type,"size_bytes":len(content),"uploaded_by_account_id":account["id"],"uploaded_by":account["name"],"uploaded_at":_now()}
    try: return _insert("source_registers",row)
    except Exception:
        if not _test_mode:
            try: supabase.storage.from_(SOURCE_REGISTERS_BUCKET).remove([path])
            except Exception: pass
        raise

@router.get("/changes")
def changes(account=Depends(_session)):
    own=[row for row in _all("changes",True) if row.get("actor_name")==account["name"]]
    return {"undo":next((row for row in own if not row.get("undone_at")),None),"redo":next((row for row in own if row.get("undone_at")),None)}

@router.post("/changes/{direction}")
def restore_change(direction:Literal["undo","redo"],account=Depends(_operator)):
    own=[row for row in _all("changes",True) if row.get("actor_name")==account["name"]]
    change=next((row for row in own if bool(row.get("undone_at"))==(direction=="redo")),None)
    if not change: raise HTTPException(409,f"Nothing to {direction}.")
    if _test_mode:
        desired=change["after_state"] if direction=="redo" else change["before_state"]; _update("tools",change["tool_id"],desired)
        _update("changes",change["id"],{"undone_at":None if direction=="redo" else _now(),"undone_by":None if direction=="redo" else account["name"]})
        return {"tool":desired,"change_id":change["id"],"redo":direction=="redo"}
    response=supabase.rpc("tools_workspace_restore_change",{"change_key":change["id"],"actor":account["name"],"redo":direction=="redo"}).execute()
    return getattr(response,"data",None)

ALIASES={"tool id":"register_number","equipment id":"register_number","register number":"register_number","employee number":"employee_number","employee no":"employee_number","name":"name","tool":"name","equipment":"name","make/model":"make_model","model":"make_model","serial number":"serial_number","serial":"serial_number","location":"storage_location","storage location":"storage_location","department":"department","job title":"job_title","category":"category"}
def _rows_from_file(content:bytes,filename:str):
    suffix=filename.lower().rsplit(".",1)[-1]
    if suffix=="csv": return list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
    if suffix in {"xlsx","xlsm"}:
        book=load_workbook(io.BytesIO(content),read_only=True,data_only=True); return [list(row) for row in book.active.iter_rows(values_only=True)]
    if suffix=="docx":
        with zipfile.ZipFile(io.BytesIO(content)) as archive: xml=archive.read("word/document.xml")
        root=ElementTree.fromstring(xml); ns={"w":"http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        return [["".join(cell.itertext()).strip() for cell in row.findall("w:tc",ns)] for row in root.findall(".//w:tr",ns)]
    raise HTTPException(415,"Choose an XLSX, XLSM, CSV or DOCX file.")
@router.post("/imports/preview")
async def preview_import(target:Literal["equipment","employees"]=Form(...),file:UploadFile=File(...),_=Depends(_session)):
    content=await file.read()
    if len(content)>20*1024*1024: raise HTTPException(413,"Import files must be smaller than 20 MB.")
    raw=[row for row in _rows_from_file(content,file.filename or "") if any(str(value or "").strip() for value in row)]
    if len(raw)<2: raise HTTPException(422,"The file needs a header row and at least one data row.")
    headers=[ALIASES.get(str(value or "").strip().lower(),str(value or "").strip().lower().replace(" ","_")) for value in raw[0]]
    required={"equipment":{"register_number","name","storage_location"},"employees":{"employee_number","name","department"}}[target]
    mapped=[{headers[i]:str(value or "").strip() for i,value in enumerate(row) if i<len(headers)} for row in raw[1:501]]
    preview=[{"row":i+2,"data":item,"errors":[f"Missing {field.replace('_',' ')}" for field in required if not item.get(field)]} for i,item in enumerate(mapped)]
    return {"target":target,"columns":headers,"rows":preview,"valid":sum(not row["errors"] for row in preview),"invalid":sum(bool(row["errors"]) for row in preview),"truncated":len(raw)>501}

@router.post("/imports/commit")
def commit_import(body:ImportCommit,account=Depends(_operator)):
    accepted=[]; rejected=[]
    for index,item in enumerate(body.rows,1):
        try:
            if body.target=="employees":
                model=EmployeeInput.model_validate(item)
                _ensure_managed_department(account,model.department)
                if _find("employees","employee_number",model.employee_number): raise ValueError("Employee number already exists")
                accepted.append(_insert("employees",{"id":str(uuid.uuid4()),**model.model_dump(),"active":True,"created_by":account["name"],"created_at":_now()}))
            else:
                model=ToolInput.model_validate(item)
                _ensure_managed_department(account,model.department)
                if _find("tools","register_number",model.register_number): raise ValueError("Register number already exists")
                accepted.append(_apply_change(None,{"id":str(uuid.uuid4()),**model.model_dump(),"status":"available","custody":None},"imported",account["name"],"Imported into the register")["tool"])
        except Exception as exc:
            rejected.append({"row":index,"error":str(exc)})
    return {"accepted":accepted,"rejected":rejected,"accepted_count":len(accepted),"rejected_count":len(rejected)}

def _reset_for_tests():
    global _test_mode; _test_mode=True
    for values in _memory.values(): values.clear()
